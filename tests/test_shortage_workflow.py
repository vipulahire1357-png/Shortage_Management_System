import os
import shutil
import sqlite3
import tempfile
import unittest
from pathlib import Path
import openpyxl
import pandas as pd

from backend.app.database import DB_COLUMNS, API_COLUMNS
from services.shortage_export import build_material_shortage_workbook
from services.vendortrack_service import (
    validate_material_shortage_register,
    read_register_to_vendor_track_rows,
    refresh_vendor_track_from_registers,
    REQUIRED_REGISTER_FIELDS,
)
from services.vendor_request_export import (
    calculate_batch_remark,
    sanitize_vendor_folder_name,
    build_vendor_request_workbook,
    export_vendor_date_request,
    VENDOR_REQUEST_COLUMNS,
)


class TestMaterialShortageColumns(unittest.TestCase):
    def test_vendor_detail_position_in_columns(self):
        """1. Vendor Detail appears immediately after Free Stock and immediately before Remarks."""
        free_idx = DB_COLUMNS.index("Free Stock")
        vendor_idx = DB_COLUMNS.index("Vendor Detail")
        remarks_idx = DB_COLUMNS.index("Remarks")
        self.assertEqual(vendor_idx, free_idx + 1)
        self.assertEqual(remarks_idx, vendor_idx + 1)

        api_free_idx = API_COLUMNS.index("free_stock")
        api_vendor_idx = API_COLUMNS.index("vendor_detail")
        api_remarks_idx = API_COLUMNS.index("remarks")
        self.assertEqual(api_vendor_idx, api_free_idx + 1)
        self.assertEqual(api_remarks_idx, api_vendor_idx + 1)


class TestShortageExport(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_material_shortage_workbook_generation(self):
        """2 & 3. Material Shortage Register export includes Vendor Detail, has proper formatting and columns."""
        rows = [
            {
                "Scope": "ELECTRICAL",
                "Material Ref Code": "REF001",
                "WBS Element": "IN-TEST-01",
                "Material Code": "0000012345",
                "Material Description": "Contactor 24V",
                "ZPP7 Quantity": 10.0,
                "Project Stock Quantity": 2.0,
                "PO Quantity": 5.0,
                "Open PO Quantity": 2.0,
                "PR Quantity": 1.0,
                "Gross Shortage": 5.0,
                "Plant Stock": 2.0,
                "Free Stock": 1.0,
                "Vendor Detail": "1100001469 DEVSHREE ENTERPRISES",
                "Remarks": "Urgent",
                "Net Shortage": 4.0,
            }
        ]

        output_path = Path(self.test_dir) / "TEST_PROJ.xlsx"
        wb = build_material_shortage_workbook(rows, table_name_suffix="TEST_PROJ")
        wb.save(output_path)

        self.assertTrue(output_path.exists())
        loaded_wb = openpyxl.load_workbook(output_path)
        self.assertEqual(loaded_wb.active.title, "Material Shortage Register")

        ws = loaded_wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Vendor Detail", headers)
        free_idx = headers.index("Free Stock")
        vend_idx = headers.index("Vendor Detail")
        rem_idx = headers.index("Remarks")
        self.assertEqual(vend_idx, free_idx + 1)
        self.assertEqual(rem_idx, vend_idx + 1)

        # Material code preserved as string
        mat_cell = ws.cell(row=2, column=headers.index("Material Code") + 1)
        self.assertEqual(str(mat_cell.value), "0000012345")


class TestVendorTrackIngestion(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_validation_detects_invalid_file(self):
        """Validate register files and detect invalid structure."""
        invalid_path = Path(self.test_dir) / "invalid.xlsx"
        df = pd.DataFrame([{"Random Col": 123}])
        df.to_excel(invalid_path, index=False)

        with self.assertRaises(ValueError) as ctx:
            validate_material_shortage_register(invalid_path)
        self.assertIn("missing required column", str(ctx.exception).lower())

    def test_read_register_to_vendor_track_rows(self):
        """5 & 8. Extract exact 8 columns and split multi-vendors."""
        valid_path = Path(self.test_dir) / "PROJ_A.xlsx"
        df = pd.DataFrame([
            {
                "WBS Element": "WBS-001",
                "Material Code": "MAT-001",
                "Material Description": "Relay",
                "Net Shortage": 12.0,
                "Vendor Detail": "1100001469 DEVSHREE ENTERPRISES / 1100009999 SIEMENS",
            }
        ])
        df.to_excel(valid_path, index=False)

        resolved = validate_material_shortage_register(valid_path)
        self.assertIsInstance(resolved, dict)

        rows = read_register_to_vendor_track_rows(valid_path, "PROJ_A")
        # Multi-vendor split into 2 rows
        self.assertEqual(len(rows), 2)
        v1 = rows[0]
        self.assertEqual(v1[0], "PROJ_A")
        self.assertEqual(v1[1], "1100001469 DEVSHREE ENTERPRISES")
        self.assertEqual(v1[2], "WBS-001")
        self.assertEqual(v1[3], "MAT-001")
        self.assertEqual(v1[4], "Relay")
        self.assertEqual(v1[5], 12.0)
        self.assertEqual(v1[6], "")
        self.assertEqual(v1[7], "")
        self.assertEqual(v1[8], "")

    def test_refresh_multi_projects_and_preserve_dates(self):
        """4, 5, 6, 7. Multiple projects combined, unselected excluded, no duplicates on re-run, dates preserved."""
        db_path = Path(self.test_dir) / "services_test.db"

        # Register for PROJ-1
        p1_path = Path(self.test_dir) / "PROJ-1.xlsx"
        df1 = pd.DataFrame([
            {
                "WBS Element": "WBS-P1",
                "Material Code": "MAT-001",
                "Material Description": "Relay",
                "Net Shortage": 5.0,
                "Vendor Detail": "1100001469 DEVSHREE ENTERPRISES",
            }
        ])
        df1.to_excel(p1_path, index=False)

        # Register for PROJ-2
        p2_path = Path(self.test_dir) / "PROJ-2.xlsx"
        df2 = pd.DataFrame([
            {
                "WBS Element": "WBS-P2",
                "Material Code": "MAT-002",
                "Material Description": "Contactor",
                "Net Shortage": 8.0,
                "Vendor Detail": "1100009999 SIEMENS",
            }
        ])
        df2.to_excel(p2_path, index=False)

        # 1. Process both projects
        res = refresh_vendor_track_from_registers(
            ["PROJ-1", "PROJ-2"], registers_dir=Path(self.test_dir), db_path=db_path
        )
        self.assertEqual(res["total_rows"], 2)
        self.assertEqual(len(res["projects"]), 2)

        # Verify DB has 2 rows (one for each project)
        conn = sqlite3.connect(db_path)
        cur = conn.cursor()
        rows = cur.execute('SELECT "project_id", "Material Code", "Vendor Date" FROM vendor_track ORDER BY "project_id"').fetchall()
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][0], "PROJ-1")
        self.assertEqual(rows[1][0], "PROJ-2")

        # 2. Update a date in PROJ-1
        conn.execute('UPDATE vendor_track SET "Vendor Date" = "2026-12-01" WHERE "project_id" = "PROJ-1"')
        conn.commit()

        # 3. Reprocess only PROJ-1: must NOT duplicate rows, PROJ-2 must remain intact, and date must be preserved
        refresh_vendor_track_from_registers(
            ["PROJ-1"], registers_dir=Path(self.test_dir), db_path=db_path
        )

        rows_after = cur.execute('SELECT "project_id", "Material Code", "Vendor Date" FROM vendor_track ORDER BY "project_id"').fetchall()
        self.assertEqual(len(rows_after), 2)  # Still 2 rows (no duplicates!)
        self.assertEqual(rows_after[0][0], "PROJ-1")
        self.assertEqual(rows_after[0][2], "2026-12-01")  # Date preserved!
        self.assertEqual(rows_after[1][0], "PROJ-2")  # PROJ-2 untouched!
        conn.close()


class TestFirstSecondRemarkRule(unittest.TestCase):
    def test_first_remark_when_all_dates_empty(self):
        """12. First is assigned when all date values are empty across all rows."""
        rows = [
            {"Vendor Date": "", "Proposed Date": None, "Final Date": "   "},
            {"Vendor Date": None, "Proposed Date": "", "Final Date": ""},
        ]
        self.assertEqual(calculate_batch_remark(rows), "First")

    def test_second_remark_when_vendor_date_populated(self):
        """13. Second is assigned when any date is populated in any row."""
        rows = [
            {"Vendor Date": "2026-11-01", "Proposed Date": "", "Final Date": ""},
            {"Vendor Date": "", "Proposed Date": "", "Final Date": ""},
        ]
        self.assertEqual(calculate_batch_remark(rows), "Second")

    def test_second_remark_when_proposed_or_final_date_populated(self):
        rows1 = [{"Vendor Date": "", "Proposed Date": "2026-11-05", "Final Date": ""}]
        self.assertEqual(calculate_batch_remark(rows1), "Second")

        rows2 = [{"Vendor Date": "", "Proposed Date": "", "Final Date": "2026-11-10"}]
        self.assertEqual(calculate_batch_remark(rows2), "Second")


class TestVendorDateRequestExport(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.db_path = Path(self.test_dir) / "test_services.db"
        # Create a test sqlite db with vendor_track table
        conn = sqlite3.connect(self.db_path)
        with conn:
            conn.execute("""
                CREATE TABLE vendor_track (
                    "project_id" TEXT,
                    "Vendor Detail" TEXT,
                    "WBS Element" TEXT,
                    "Material Code" TEXT,
                    "Description" TEXT,
                    "Quantity" REAL,
                    "Vendor Date" TEXT,
                    "Proposed Date" TEXT,
                    "Final Date" TEXT,
                    "Status" TEXT
                )
            """)
        conn.close()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_sanitize_vendor_folder_name(self):
        self.assertEqual(
            sanitize_vendor_folder_name("1100001469 DEVSHREE / ENTERPRISES: *?"),
            "1100001469 DEVSHREE ENTERPRISES"
        )

    def test_export_rejects_all_vendors(self):
        """14. Export rejects All Vendors."""
        with self.assertRaises(ValueError) as ctx:
            export_vendor_date_request(
                vendor="All Vendors",
                db_path=self.db_path,
                output_root=self.test_dir,
            )
        self.assertIn("select a specific vendor", str(ctx.exception))

    def test_export_rejects_empty_results(self):
        """14. Export rejects empty results."""
        with self.assertRaises(ValueError) as ctx:
            export_vendor_date_request(
                vendor="1100001469 DEVSHREE ENTERPRISES",
                db_path=self.db_path,
                output_root=self.test_dir,
            )
        self.assertIn("No records found", str(ctx.exception))

    def test_export_creates_folder_and_file_with_versioning(self):
        """9, 15, 16. Exports to Vendor Registers/<vendor>/Outgoing/Vendor_Date_Request/{First,Second}.xlsx with versioning."""
        # Insert records with empty dates
        conn = sqlite3.connect(self.db_path)
        with conn:
            conn.execute(
                """
                INSERT INTO vendor_track
                ("project_id", "Vendor Detail", "WBS Element", "Material Code", "Description", "Quantity", "Vendor Date", "Proposed Date", "Final Date")
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                ("PROJ-1", "1100001469 DEVSHREE ENTERPRISES", "IN-TEST-01", "000123", "Relay", 10.0, "", "", ""),
            )
        conn.close()

        # First export -> First.xlsx
        res1 = export_vendor_date_request(
            vendor="1100001469 DEVSHREE ENTERPRISES",
            db_path=self.db_path,
            output_root=self.test_dir,
        )
        self.assertEqual(res1["remark"], "First")
        self.assertEqual(res1["filename"], "First.xlsx")
        self.assertTrue(Path(res1["output_path"]).exists())
        self.assertIn("Outgoing", res1["output_path"])
        self.assertIn("Vendor_Date_Request", res1["output_path"])

        # Second export of First condition should create First_v2.xlsx
        res2 = export_vendor_date_request(
            vendor="1100001469 DEVSHREE ENTERPRISES",
            db_path=self.db_path,
            output_root=self.test_dir,
        )
        self.assertEqual(res2["remark"], "First")
        self.assertEqual(res2["filename"], "First_v2.xlsx")
        self.assertTrue(Path(res2["output_path"]).exists())

        # Now update row with a date -> Second.xlsx
        conn = sqlite3.connect(self.db_path)
        with conn:
            conn.execute(
                """
                UPDATE vendor_track
                SET "Vendor Date" = '2026-11-20'
                WHERE "Vendor Detail" = '1100001469 DEVSHREE ENTERPRISES'
                """
            )
        conn.close()

        res3 = export_vendor_date_request(
            vendor="1100001469 DEVSHREE ENTERPRISES",
            db_path=self.db_path,
            output_root=self.test_dir,
        )
        self.assertEqual(res3["remark"], "Second")
        self.assertEqual(res3["filename"], "Second.xlsx")
        self.assertTrue(Path(res3["output_path"]).exists())

        # Verify Excel content has 8 exact columns in order
        wb = openpyxl.load_workbook(res3["output_path"])
        ws = wb.active
        self.assertEqual(ws.title, "Vendor Date Request")
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(headers, VENDOR_REQUEST_COLUMNS)

        # Verify Material Code preserved as text
        mat_cell = ws.cell(row=2, column=headers.index("Material Code") + 1)
        self.assertEqual(str(mat_cell.value), "000123")


class TestME2MReportDiscovery(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_reports_definition_has_me2m(self):
        """REPORTS in database.py contains ME2M with label, stem, and .xlsb support."""
        from backend.app.database import REPORTS
        self.assertIn("me2m", REPORTS)
        me2m_meta = REPORTS["me2m"]
        self.assertEqual(me2m_meta["stem"], "ME2M")
        self.assertIn(".xlsb", me2m_meta["extensions"])
        self.assertIn("Vendor / Material Details", me2m_meta["label"])

    def test_me2m_classifier_and_discovery(self):
        """_is_me2m classifies ME2M correctly and does not match other reports."""
        from backend.app.projects import _is_me2m, _classify_files, PROJECT_REPORT_CLASSIFIERS

        # Matches valid ME2M files (including .xlsb)
        self.assertTrue(_is_me2m(Path("ME2M  aLL  DATA_ (version 1).xlsb")))
        self.assertTrue(_is_me2m(Path("me2m_IN-10026019.XLSX")))
        self.assertTrue(_is_me2m(Path("ME2M.xls")))

        # Does NOT match other reports
        self.assertFalse(_is_me2m(Path("me2j_IN-10026019.XLSX")))
        self.assertFalse(_is_me2m(Path("mb25_IN-10026019.XLSX")))
        self.assertFalse(_is_me2m(Path("zpp7_IN-10026019.xls")))
        self.assertFalse(_is_me2m(Path("pr_IN-10026019.XLSX")))

        # Folder classification
        folder = Path(self.test_dir)
        (folder / "ME2M  aLL  DATA_ (version 1).xlsb").touch()
        (folder / "zpp7_test.xls").touch()
        (folder / "mb25_test.xlsx").touch()
        (folder / "me2j_test.xlsx").touch()
        (folder / "pr_test.xlsx").touch()

        matches = _classify_files(folder, PROJECT_REPORT_CLASSIFIERS)
        self.assertIsNotNone(matches.get("me2m"))
        self.assertEqual(matches["me2m"].name, "ME2M  aLL  DATA_ (version 1).xlsb")

    def test_missing_me2m_returns_none(self):
        """Folder without ME2M classifies me2m as None."""
        from backend.app.projects import _classify_files, PROJECT_REPORT_CLASSIFIERS

        folder = Path(self.test_dir)
        (folder / "zpp7_test.xls").touch()

        matches = _classify_files(folder, PROJECT_REPORT_CLASSIFIERS)
        self.assertIsNone(matches.get("me2m"))


if __name__ == "__main__":
    unittest.main()
